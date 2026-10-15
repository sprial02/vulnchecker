"""Client-facing DOCX report with customer footer and streamed complete CSV appendix."""
import argparse
import csv
import hashlib
from datetime import datetime
import importlib.util
import json
from pathlib import Path
import re
import tempfile
import uuid
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

spec = importlib.util.spec_from_file_location('report_store', Path(__file__).with_name('report-store.py'))
store = importlib.util.module_from_spec(spec); spec.loader.exec_module(store)
matrix_spec=importlib.util.spec_from_file_location('report_matrix',Path(__file__).with_name('report-matrix.py'))
matrix=importlib.util.module_from_spec(matrix_spec);matrix_spec.loader.exec_module(matrix)
TYPES = {'mobile': '모바일 앱', 'web': '웹서버', 'infra': '인프라 서버 및 네트워크'}
LABELS = {'GOOD':'양호','VULNERABLE':'취약','MANUAL':'수동 확인','N/A':'해당 없음',
          'Critical':'긴급','High':'높음','Medium':'중간','Low':'낮음','Info':'정보','Unknown':'미분류'}
RISK_ORDER = "CASE severity WHEN 'Critical' THEN 0 WHEN 'High' THEN 1 WHEN 'Medium' THEN 2 WHEN 'Low' THEN 3 ELSE 4 END"


def font(run, size=10.5, bold=False, color='000000'):
    run.font.name = '맑은 고딕'; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    return run


def paragraph(doc, value='', style=None, bold=False, size=10.5):
    p = doc.add_paragraph(style=style); font(p.add_run(store.text(value)), size, bold)
    return p


def field(p, name):
    node = OxmlElement('w:fldSimple'); node.set(qn('w:instr'), name)
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '1'; run.append(text); node.append(run); p._p.append(node)


def setup(customer):
    doc = Document(); section = doc.sections[0]
    section.page_width = Cm(21.59); section.page_height = Cm(27.94)
    section.top_margin = Cm(2.0); section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.395); section.right_margin = Cm(2.395)
    section.header_distance = Cm(0.9); section.footer_distance = Cm(0.9)
    for name, size in [('Normal',10.5),('Title',27),('Subtitle',14),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
        style = doc.styles[name]; style.font.name='맑은 고딕'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
        style.paragraph_format.space_after=Pt(7); style.paragraph_format.line_spacing=1.2
        if name.startswith('Heading'): style.paragraph_format.keep_with_next=True; style.paragraph_format.space_before=Pt(12)
    title_properties = doc.styles['Title']._element.get_or_add_pPr()
    for node in list(title_properties):
        if node.tag == qn('w:pBdr'): title_properties.remove(node)
    header = section.header.paragraphs[0]
    font(header.add_run('IBM TriGuard AI'), 8)
    footer = section.footer.paragraphs[0]
    doc.styles['Footer'].paragraph_format.tab_stops.clear_all()
    footer.paragraph_format.tab_stops.clear_all()
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), WD_TAB_ALIGNMENT.RIGHT)
    font(footer.add_run(customer), 8)
    font(footer.add_run('\t대외비  |  '), 8)
    field(footer,'PAGE'); font(footer.add_run(' / '),8); field(footer,'NUMPAGES')
    doc.core_properties.author = 'IBM TriGuard AI'
    return doc


def heading(doc, title, level=1):
    p=doc.add_paragraph(title, 'Heading '+str(level))
    for run in p.runs: font(run, {1:17,2:13,3:11}[level], True)
    return p


def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for column,width in zip(t.columns,widths): column.width=Cm(width)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        node=OxmlElement('w:'+edge); node.set(qn('w:val'),'single'); node.set(qn('w:sz'),'4'); node.set(qn('w:color'),'D9D9D9'); borders.append(node)
    t._tbl.tblPr.append(borders)
    for index, row in enumerate([headers]+list(rows)):
        cells=t.rows[0].cells if index==0 else t.add_row().cells
        for i,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'24374B' if index==0 else ('F2F5F8' if index%2==0 else 'FFFFFF')); props.append(shade)
            margins=OxmlElement('w:tcMar')
            for name,value_margin in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
                n=OxmlElement('w:'+name); n.set(qn('w:w'),value_margin); n.set(qn('w:type'),'dxa'); margins.append(n)
            props.append(margins)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if widths[i]<3 else WD_ALIGN_PARAGRAPH.LEFT
            font(p.add_run(store.text(value,5000)),9,index==0,'FFFFFF' if index==0 else '000000')
        if index==0:
            repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    paragraph(doc,'').paragraph_format.space_after=Pt(0)
    return t


def chart(path, counts):
    image=Image.new('RGB',(1200,380),'white'); draw=ImageDraw.Draw(image)
    font_path=Path('C:/Windows/Fonts/malgun.ttf')
    f=ImageFont.truetype(str(font_path),27) if font_path.exists() else ImageFont.load_default()
    colors={'GOOD':'#4D7A68','VULNERABLE':'#A74743','MANUAL':'#B28B37','N/A':'#8A94A0'}
    maximum=max([1]+list(counts.values()))
    for i,status in enumerate(colors):
        y=15+i*87; count=counts.get(status,0)
        draw.text((10,y+12),LABELS[status],font=f,fill='black')
        width=int(780*count/maximum); draw.rectangle((235,y+10,235+max(2,width),y+53),fill=colors[status])
        draw.text((1040,y+12),f'{count:,}',font=f,fill='black')
    image.save(path)


def csv_cell(value):
    value=store.text(value,100000)
    return "'"+value if value[:1] in ('=','+','-','@') else value


def export_csv(db, source, path):
    with path.open('w',encoding='utf-8-sig',newline='') as file:
        writer=csv.writer(file); writer.writerow(['자산','항목 ID','항목명','결과','위험도','분류','상세 내용','영향','조치 권고','증적','근거','원본 위치'])
        for row in db.execute('SELECT asset,item_id,title,status,severity,category,description,impact,recommendation,evidence,reference,origin FROM checks WHERE source_id=? ORDER BY asset,item_id',(source,)):
            values=list(row); values[3]=LABELS.get(values[3],values[3]); values[4]=LABELS.get(values[4],values[4]); writer.writerow([csv_cell(v) for v in values])


def generate(db, request, out_dir, template=False):
    customer=str(request.get('Customer') or '고객사명').strip()
    if not customer or '\n' in customer or len(customer)>100: raise ValueError('고객사명은 1~100자로 입력하세요.')
    kind=request.get('Kind','infra')
    if kind not in TYPES: raise ValueError('보고서 종류가 올바르지 않습니다.')
    source_id=request.get('SourceId')
    source=db.execute('SELECT * FROM sources WHERE id=?',(source_id,)).fetchone() if source_id and not template else None
    if source_id and not source and not template: raise ValueError('선택한 결과를 찾을 수 없습니다.')
    if source and source['kind']!=kind: raise ValueError('선택한 결과 종류가 보고서와 다릅니다.')
    if not source and not template: raise ValueError('점검 결과를 선택하세요. 결과 없이 출력하려면 3종 서식 추출을 사용하세요.')
    stats=store.stats(db,source_id) if source else None
    draft=not source or not stats['total_checks'] or stats['failed_assets']>0 or stats['detailed_assets']<stats['assets'] or not all(request.get(k) for k in ('Author','Period','Scope'))
    if kind=='infra': return generate_infra(db,request,out_dir,template,source,stats,draft)
    doc=setup(customer)
    title=f'{customer} {TYPES[kind]} 취약점 점검 보고서'
    doc.core_properties.title=title
    paragraph(doc,'보안 점검 보고서',bold=True,size=11)
    paragraph(doc,'IBM TriGuard AI',bold=True,size=16)
    paragraph(doc,'Integrated Automated Vulnerability Assessment & Reporting Suite (Korea Delivery)',size=9)
    paragraph(doc,title,style='Title',size=27,bold=True)
    paragraph(doc,'검토용 서식' if template else ('결과 검토본' if draft else '점검 결과 보고서'),style='Subtitle',size=14)
    paragraph(doc,request.get('Project') or TYPES[kind]+' 취약점 점검',size=13)
    paragraph(doc,'작성일  '+datetime.now().strftime('%Y년 %m월 %d일'))
    paragraph(doc,'점검 기간  '+str(request.get('Period') or '실제 점검 기간 입력'))
    paragraph(doc,'작성 및 검토 담당  '+str(request.get('Author') or '담당자 입력'))
    paragraph(doc,'문서 버전  1.0')
    paragraph(doc,'본 문서는 승인된 점검 범위의 결과를 경영진과 운영 담당자가 함께 검토하고 개선 우선순위를 결정하기 위한 자료입니다. 상세 증적에는 내부 구성 정보가 포함될 수 있으므로 고객사의 문서 관리 정책에 따라 취급하십시오.')
    if template: paragraph(doc,'이 문서는 보고서 서식입니다. 실측 점검 결과와 담당자 검토 내용을 입력한 뒤 고객 제출본으로 확정하십시오. 아래 빈 통계는 양호 또는 취약 여부를 의미하지 않습니다.',bold=True)
    doc.add_page_break()
    heading(doc,'경영진 요약 Executive Summary')
    if template:
        paragraph(doc,'핵심 결론과 사업 영향을 작성합니다. 실제 점검 결과에 근거하여 우선 조치할 위험, 운영상 제약, 경영진이 결정해야 할 사항을 설명하십시오. 미확인 영역은 판단이 유보되어 있음을 명시합니다.')
    else:
        counts=stats['counts']; vulnerabilities=counts.get('VULNERABLE',0)
        paragraph(doc,f"총 {stats['assets']:,}개 자산에서 {stats['total_checks']:,}건의 항목 결과가 집계되었습니다. 취약 판정은 {vulnerabilities:,}건이며 {stats['affected_assets']:,}개 자산에 해당합니다. 수동 확인은 {counts.get('MANUAL',0):,}건입니다. 같은 취약점이 여러 자산에 나타나면 자산별 판정 건수로 집계합니다.")
        if draft: paragraph(doc,f"점검 수행 실패 자산은 {stats['failed_assets']:,}개이며 상세 증적이 요약 합계와 일치하는 자산은 {stats['detailed_assets']:,}/{stats['assets']:,}개입니다. 결과가 완전하지 않은 영역을 재점검하고 이 보고서의 결론을 확정해야 합니다.",bold=True)
        if vulnerabilities and sum(stats['severity'].get(s,0) for s in ('Critical','High')):
            paragraph(doc,'긴급 또는 높은 위험도로 기록된 항목을 우선 검토하십시오. 서비스 영향과 변경 위험을 평가한 후 임시 완화 및 영구 조치 계획을 결정해야 합니다.')
        elif vulnerabilities:
            paragraph(doc,'취약 판정 항목의 영향 범위와 위험도를 검토해 조치 순서를 확정하십시오. 위험도가 입력되지 않은 항목은 낮은 위험으로 간주하지 않고 별도 검토 대상으로 유지했습니다.')
        else:
            paragraph(doc,'제공된 결과에서 취약 판정이 집계되지 않았습니다. 이 사실만으로 미점검 자산이나 수동 확인 영역의 안전성을 보장할 수는 없습니다.')
    if request.get('ExecutiveNote'): paragraph(doc,request['ExecutiveNote'])
    keys=[('점검 대상 자산','assets'),('수행 실패 자산','failed_assets'),('취약 자산','affected_assets'),('항목 결과 합계','total_checks')]
    kpis=[(label,f'{stats[key]:,}' if stats else '미입력') for label,key in keys]
    kpis += [(LABELS[status]+' 판정',f"{stats['counts'].get(status,0):,}" if stats else '미입력') for status in ('GOOD','VULNERABLE','MANUAL','N/A')]
    table(doc,['지표','집계 값','지표','집계 값'],[kpis[i]+kpis[i+1] for i in range(0,len(kpis),2)],[5.4,3.0,5.4,3.0])
    paragraph(doc,'항목 결과 합계는 자산별 점검 결과의 합입니다. 양호율은 양호 ÷ 양호와 취약의 합으로 계산하며 수동 확인과 해당 없음은 분모에서 제외합니다. 수행 실패와 상세 증적 누락은 별도 표시합니다.',size=9)
    if stats:
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'status.png'; chart(path,stats['counts']); doc.add_picture(str(path),width=Cm(16.5))
        paragraph(doc,'자산별 판정 건수 분포',size=9)
    heading(doc,'우선 의사결정과 권고',2)
    paragraph(doc,'우선 조치 항목의 담당 조직과 목표 일정을 지정하십시오. 시스템 변경 전 백업과 복구 방법을 확인하고, 변경 후 재점검에서 동일 항목의 개선 여부와 서비스 정상 동작을 확인해야 합니다.')
    doc.add_page_break()
    heading(doc,'점검 개요와 결과 분석')
    heading(doc,'대상과 수행 범위',2)
    paragraph(doc,request.get('Scope') or ('실제 점검한 앱 패키지 버전과 API 범위를 입력합니다.' if kind=='mobile' else '실제 점검한 서비스 주소와 역할 및 제외 범위를 입력합니다.' if kind=='web' else '실제 점검한 서버와 네트워크 장비의 범위 및 제외 대상을 입력합니다.'))
    paragraph(doc,'결과는 제공된 원본과 수집된 증적 범위에서 작성합니다. 계정 권한 부족, 연결 실패, 수동 판단 필요 항목과 제외 범위를 기록하고 추가 확인 결과에 따라 결론을 갱신합니다.')
    if source:
        paragraph(doc,'결과 식별자  '+source['name'],size=9)
        metadata=json.loads(source['metadata'])
        if isinstance(metadata,dict) and metadata.get('scope'): paragraph(doc,'원본 수행 범위  '+store.text(metadata['scope']))
        if isinstance(metadata,dict) and metadata.get('limitations'): paragraph(doc,'수행 제약과 제외 범위  '+store.text(metadata['limitations']))
        if metadata.get('revision'): paragraph(doc,'점검 기준 소스 리비전  '+str(metadata['revision']),size=9)
    heading(doc,'위험도별 취약 판정',2)
    risks=[(LABELS[s],f"{stats['severity'].get(s,0):,}" if stats else '미입력') for s in ('Critical','High','Medium','Low','Info','Unknown')]
    table(doc,['위험도','상세 증적에서 확인된 건수'],risks,[5.0,11.8])
    paragraph(doc,'위험도는 원본 결과의 값을 사용합니다. 위험도가 없으면 미분류로 표시하며 자동으로 CVSS 점수나 위험도를 만들어 넣지 않습니다. 상세 증적이 부족하면 위 위험도 합계가 전체 취약 판정 합계와 다를 수 있습니다.',size=9)
    if stats:
        heading(doc,'자산별 집계와 확인 범위',2)
        assets=db.execute('''SELECT a.name,a.platform,a.failed,a.detailed,
          COALESCE(SUM(CASE WHEN r.status='VULNERABLE' THEN r.count ELSE 0 END),0) AS vuln,
          COALESCE(SUM(r.count),0) AS total FROM assets a LEFT JOIN result_counts r
          ON r.source_id=a.source_id AND r.asset=a.name WHERE a.source_id=?
          GROUP BY a.name ORDER BY vuln DESC,a.name LIMIT 20''',(source_id,)).fetchall()
        table(doc,['자산','유형','전체','취약','증적'],[(a['name'],a['platform'],a['total'],a['vuln'],'실패' if a['failed'] else ('확인' if a['detailed'] else '미완전')) for a in assets],[6.0,3.1,1.8,1.8,4.1])
        paragraph(doc,'본문 자산 집계는 취약 판정이 많은 상위 20개 자산입니다. 전체 자산 집계는 별도 자산 CSV 부록에서 확인합니다.',size=9)
    doc.add_page_break()
    heading(doc,'주요 취약점 상세와 개선 계획')
    limit=max(1,min(200,int(request.get('DetailLimit',30))))
    details=db.execute(f"SELECT * FROM checks WHERE source_id=? AND status IN ('VULNERABLE','MANUAL') ORDER BY CASE status WHEN 'VULNERABLE' THEN 0 ELSE 1 END,{RISK_ORDER},asset,item_id LIMIT ?",(source_id,limit)).fetchall() if source else []
    if not details:
        heading(doc,'취약점 사례 작성',2)
        paragraph(doc,'확인된 결과가 없으면 취약점 사례를 작성하지 않습니다. 아래 항목은 실측 결과를 검토할 때 채워야 하는 구성입니다.' if template else '검증 가능한 취약 또는 수동 확인 상세 항목이 없습니다. 요약 판정에 취약 건수가 있으면 원본 증적을 추가 수집하십시오.')
        if template:
            for label in ('항목 식별자와 취약점명','영향받는 자산과 위험도','관찰 내용과 재현 조건','기술적 영향과 사업 영향','증적과 판정 근거','권고 조치와 검증 방법'):
                heading(doc,label,3); paragraph(doc,'실제 점검 결과와 담당자 검토 내용을 입력합니다.')
    for index,row in enumerate(details,1):
        if index>1: doc.add_page_break()
        heading(doc,f'{index} {row["item_id"]} {row["title"]}',2)
        paragraph(doc,f"자산 {row['asset']}   결과 {LABELS[row['status']]}   위험도 {LABELS[row['severity']]}",bold=True,size=10)
        paragraph(doc,row['description'] or '원본의 상세 판정 사유를 검토해 보완하십시오.')
        heading(doc,'영향',3); paragraph(doc,row['impact'] or '구체적인 영향은 원본에 기록되지 않았습니다. 서비스 역할과 노출 범위를 확인해 담당자가 평가하십시오.')
        heading(doc,'증적과 판정 근거',3)
        paragraph(doc,row['evidence'][:2500] or '상세 증적은 원본 위치를 확인하십시오.',size=9)
        paragraph(doc,'원본 위치  '+row['origin'],size=8)
        if request.get('IncludeImages'):
            for image in db.execute('SELECT path,caption,sha256 FROM attachments WHERE source_id=? AND asset=? AND item_id=?',(source_id,row['asset'],row['item_id'])):
                path=Path(image['path'])
                if not path.is_file() or path.stat().st_size>10*1024*1024 or hashlib.sha256(path.read_bytes()).hexdigest()!=image['sha256']:
                    paragraph(doc,'증적 이미지가 누락되었거나 가져온 이후 변경되어 포함하지 않았습니다.',size=9); continue
                with Image.open(path) as picture:
                    width,height=picture.size
                scale=min(16.0/width,10.0/height)
                caption=paragraph(doc,image['caption'],size=9); caption.paragraph_format.keep_with_next=True
                doc.add_picture(str(path),width=Cm(width*scale),height=Cm(height*scale))
        if row['reference']: paragraph(doc,'판정 근거  '+row['reference'],size=9)
        heading(doc,'조치 권고',3); paragraph(doc,row['recommendation'] or '원본 점검 기준과 서비스 운영 조건을 확인하여 구체적인 조치를 확정하십시오.')
    doc.add_page_break()
    heading(doc,'개선 추진과 재점검',1)
    table(doc,['단계','실행 내용','담당 및 일정'],[
      ('우선 조치','긴급 및 높은 위험과 미분류 취약점을 검토하고 임시 완화 여부 결정','담당자와 목표일 입력'),
      ('개선 실행','승인된 변경 계획에 따라 설정 수정 패치 권한 정비 수행','담당자와 목표일 입력'),
      ('재점검','동일 항목의 개선 증적과 서비스 정상 동작 확인','담당자와 목표일 입력')],[3.0,9.0,4.8])
    heading(doc,'부록과 증적 관리',1)
    paragraph(doc,f'본문 상세 사례는 최대 {limit}건으로 제한합니다. 전체 결과 CSV 부록은 이 제한과 관계없이 검증 가능한 모든 항목을 포함합니다. 원본 ZIP 또는 TAR 증적은 별도로 보존하고 항목 식별자와 원본 위치로 추적합니다.')
    paragraph(doc,'민감정보는 문서와 CSV에서 마스킹합니다. 자동 마스킹이 모든 형태의 비밀정보를 찾아내는 것은 아니므로 제출 전에 증적과 첨부 자료를 검토하십시오.',size=9)
    out_dir=Path(out_dir); out_dir.mkdir(parents=True,exist_ok=True)
    safe=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',title).strip(' .')[:120]
    base=safe+('_서식' if template else '')+'_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:6]
    doc_path=out_dir/(base+'.docx'); doc.save(doc_path)
    files=[str(doc_path)]
    if source:
        csv_path=out_dir/(base+'_전체결과.csv'); export_csv(db,source_id,csv_path); files.append(str(csv_path))
        asset_path=out_dir/(base+'_자산집계.csv')
        with asset_path.open('w',encoding='utf-8-sig',newline='') as file:
            writer=csv.writer(file); writer.writerow(['자산','유형','수행 실패','상세 증적 확인','전체','취약'])
            for row in db.execute('''SELECT a.name,a.platform,a.failed,a.detailed,COALESCE(SUM(r.count),0),
             COALESCE(SUM(CASE WHEN r.status='VULNERABLE' THEN r.count ELSE 0 END),0)
             FROM assets a LEFT JOIN result_counts r ON r.source_id=a.source_id AND r.asset=a.name
             WHERE a.source_id=? GROUP BY a.name ORDER BY a.name''',(source_id,)):
                writer.writerow([csv_cell(v) for v in row])
        files.append(str(asset_path))
    return {'Kind':kind,'Title':title,'Files':files,'Draft':draft,'Stats':stats}


def generate_infra(db,request,out_dir,template,source,stats,draft):
    customer=str(request.get('Customer') or '고객사명').strip();source_id=source['id'] if source else None
    title=customer+' 인프라 서버 및 네트워크 점검 결과 보고서'
    doc=setup(customer);doc.core_properties.title=title
    paragraph(doc,'IBM TriGuard AI',bold=True,size=16)
    paragraph(doc,'Integrated Automated Vulnerability Assessment & Reporting Suite (Korea Delivery)',size=9)
    paragraph(doc,title,style='Title',bold=True,size=27)
    paragraph(doc,'검토용 서식' if template else ('현황 검토본' if draft else '점검 결과 현황'),style='Subtitle',size=14)
    paragraph(doc,'작성일  '+datetime.now().strftime('%Y년 %m월 %d일'))
    paragraph(doc,'고객사  '+customer)
    paragraph(doc,'점검 기간  '+str(request.get('Period') or '실제 점검 기간 입력'))
    paragraph(doc,'작성 및 검토 담당  '+str(request.get('Author') or '담당자 입력'))
    paragraph(doc,'본 보고서는 전체 인프라 점검 판정과 서버별 수행 현황을 표로 정리한 결과 보고서입니다. 경영진은 전체 통계와 미완료 범위를 확인하고 운영 담당자는 서버별 항목 결과표에서 후속 확인 대상을 조회합니다. 개별 취약점의 상세 증적은 원본 자료를 참조합니다.')
    if template:paragraph(doc,'실제 결과가 입력되지 않은 보고서 서식입니다. 빈 통계와 빈 결과표는 양호 판정을 의미하지 않습니다.',bold=True)
    heading(doc,'전체 점검 결과 요약 Executive Summary').paragraph_format.page_break_before=True
    if stats:
        counts=stats['counts']
        paragraph(doc,f"물리 서버 {stats['physical_assets']:,}대, 점검 대상 {stats['assets']:,}개(OS·서비스 인스턴스)에서 {stats['total_checks']:,}건의 판정이 집계되었습니다. 취약 판정은 {counts.get('VULNERABLE',0):,}건이며 취약 판정이 있는 서버는 {stats['physical_affected_assets']:,}대입니다. 수행 실패 대상은 {stats['failed_assets']:,}개입니다. 집계 단위는 서버·점검 종류·인스턴스·항목의 조합으로, 같은 항목이 여러 서버에 나타나면 각각 계산합니다.")
        table(doc,['지표','집계 값'],[(label,stats[key]) for label,key in [('물리 서버','physical_assets'),('OS / 서비스 점검 대상','assets'),('취약 판정 물리 서버','physical_affected_assets'),('수행 실패 대상','failed_assets'),('요약과 상세 일치 대상','detailed_assets')]],[11.8,5.0])
    else:
        paragraph(doc,'실제 점검 결과의 총 서버 수, 전체 판정 합계와 수행 실패 범위를 입력합니다. 판정이 수집되지 않은 서버를 양호로 계산하지 않습니다.')
        table(doc,['지표','집계 값'],[(label,'미입력') for label in ('물리 서버','OS / 서비스 점검 대상','취약 판정 물리 서버','수행 실패 대상','요약과 상세 일치 대상')],[11.8,5.0])
    result_rows=[]
    for status in ('GOOD','VULNERABLE','MANUAL','N/A'):
        count=stats['counts'].get(status,0) if stats else None
        percentage=f"{100*count/stats['total_checks']:.1f}%" if stats and stats['total_checks'] else '산출 불가' if stats else '미입력'
        result_rows.append((LABELS[status],f'{count:,}' if count is not None else '미입력',percentage))
    result_rows.append(('판정 합계',f"{stats['total_checks']:,}" if stats else '미입력','100.0%' if stats and stats['total_checks'] else '미입력'))
    table(doc,['판정','건수','전체 판정 중 비율'],result_rows,[5.0,5.0,6.8])
    rate=str(stats['good_rate'])+'%' if stats and stats['good_rate'] is not None else '산출 불가' if stats else '미입력'
    paragraph(doc,'양호율  '+rate+'  · 양호 ÷ (양호 + 취약). 수동 확인과 해당 없음은 양호율의 분모에서 제외합니다. 위 판정 비율의 분모는 모든 판정 합계입니다.',size=9)
    if stats:
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'status.png';chart(path,stats['counts']);doc.add_picture(str(path),width=Cm(14))
    heading(doc,'점검 범위와 집계 기준').paragraph_format.page_break_before=True
    paragraph(doc,request.get('Scope') or '실제 점검한 서버 및 네트워크 장비와 제외 범위를 입력합니다.')
    if source:
        paragraph(doc,'결과 식별자  '+source['name'],size=9)
        metadata=json.loads(source['metadata'])
        for key,label in [('scope','원본 점검 범위'),('limitations','수행 제약'),('revision','점검 기준 리비전')]:
            if metadata.get(key):paragraph(doc,label+'  '+store.text(metadata[key]),size=9)
    paragraph(doc,'전체 판정과 서버별 합계는 원본 요약을 사용합니다. 점검 항목별 통계와 서버별 항목 결과표는 실제 수집된 상세 판정만 사용합니다. 요약과 상세가 일치하지 않으면 두 표의 합계가 다를 수 있으며, 이를 미완전 결과로 표시합니다. 접속 확인만 수행한 작업은 포함하지 않습니다.')
    paragraph(doc,'Excel 결과표는 서버별 OS·서비스 인스턴스가 행, 점검 종류별 항목이 열입니다. 서버명과 헤더를 고정하고 판정을 색상으로 구분합니다. 결과 없음은 상세 판정이 수집되지 않았거나 해당 서버의 점검 범위 밖인 경우이며, 해당 없음 판정을 대신하지 않습니다. 전체 결과 CSV와 원본 증적에서 판정 근거를 확인할 수 있습니다.')
    heading(doc,'점검 항목별 판정 통계').paragraph_format.page_break_before=True
    items=matrix.item_rows(db,source_id) if source else []
    rows=([matrix.profile_label(r['platform'])]+[r[k] for k in ('item_id','title','good','vulnerable','manual','na','total')] for r in items)
    table(doc,['점검 종류','항목 ID','점검 항목','양호','취약','수동','해당 없음','합계'],rows if source else [('미입력','미입력','실제 점검 항목','미입력','미입력','미입력','미입력','미입력')],[2.2,1.8,4.8,1.4,1.4,1.6,1.8,1.8])
    paragraph(doc,'각 행은 동일 점검 종류와 항목의 인스턴스별 판정 합계입니다. DBMS나 WAS 제품이 다르면 항목 ID가 같아도 분리합니다. 상세 결과가 없는 항목은 임의로 생성하지 않습니다.',size=9)
    heading(doc,'서버별 점검 결과 현황').paragraph_format.page_break_before=True
    assets=matrix.asset_rows(db,source_id) if source else []
    def asset_values():
        for r in assets:yield [(r['server_name']+' / '+r['instance_id']) if r['instance_id'] and r['instance_id']!='os' else r['name'],matrix.profile_label(r['platform']),'실패' if r['failed'] else '확인' if r['detailed'] else '미완전',r['good'],r['vulnerable'],r['manual'],r['na'],r['total']]
    table(doc,['서버 / 인스턴스','유형','수행','양호','취약','수동','해당 없음','합계'],asset_values() if source else [('미입력','미입력','미입력','미입력','미입력','미입력','미입력','미입력')],[5.0,2.0,1.8,1.4,1.4,1.6,1.8,1.8])
    paragraph(doc,'이 표는 모든 OS·서비스 점검 대상을 포함하며 같은 서버의 인스턴스는 별도 행으로 표시합니다. 서버별 항목 판정 전체는 함께 추출한 Excel 결과표에서 조회합니다. 개별 증적이나 항목별 장문의 설명을 본문에 반복하지 않습니다.',size=9)
    for t in doc.tables:
        for row in t.rows:
            node=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(node)
    out_dir=Path(out_dir);out_dir.mkdir(parents=True,exist_ok=True)
    safe=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',title).strip(' .')[:120]
    base=safe+('_서식' if template else '')+'_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:6]
    path=out_dir/(base+'.docx');doc.save(path);files=[str(path)]
    if source:
        full=out_dir/(base+'_전체결과.csv');export_csv(db,source_id,full);files.append(str(full))
        asset=out_dir/(base+'_서버집계.csv')
        with asset.open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.writer(f);writer.writerow(['서버','유형','수행 실패','상세 증적 확인','양호','취약','수동 확인','해당 없음','합계'])
            for r in matrix.asset_rows(db,source_id):writer.writerow([csv_cell(r[k]) for k in ('name','platform','failed','detailed','good','vulnerable','manual','na','total')])
        files.append(str(asset));files.append(matrix.export(db,source_id,out_dir/(base+'_서버별항목결과.xlsx'),customer,stats,source['name']))
    return {'Kind':'infra','Title':title,'Files':files,'Draft':draft,'Stats':stats}


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--db',required=True); parser.add_argument('--request',required=True); parser.add_argument('--output',required=True); parser.add_argument('--manifest',required=True); parser.add_argument('--templates',action='store_true')
    args=parser.parse_args(); request=json.loads(Path(args.request).read_text(encoding='utf-8-sig'))
    db=store.connect(args.db); db.execute('BEGIN')
    reports=[]
    for kind in TYPES if args.templates else [request['Kind']]:
        reports.append(generate(db,dict(request,Kind=kind),args.output,args.templates))
    db.rollback(); db.close()
    Path(args.manifest).write_text(json.dumps({'Reports':reports},ensure_ascii=False,indent=2),encoding='utf-8')
    print('DOCX 보고서 생성 완료')


if __name__=='__main__': main()
