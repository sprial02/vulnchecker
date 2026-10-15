"""Report data integrity and DOCX/CSV contract tests. Uses isolated temporary data."""
import csv
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from docx import Document
from PIL import Image
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
def module(name, file):
    spec=importlib.util.spec_from_file_location(name,ROOT/'scripts'/file)
    out=importlib.util.module_from_spec(spec);spec.loader.exec_module(out);return out
store=module('store','report-store.py')
report=module('report','report-docx.py')

class Reports(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name);self.db=store.connect(self.path/'results.sqlite3')
    def tearDown(self):
        self.db.close();self.tmp.cleanup()
    def write(self, rows, kind='infra'):
        path=self.path/'findings.json';path.write_text(json.dumps({'kind':kind,'name':'QA fixture','findings':rows}),encoding='utf-8');return path
    def test_failed_and_missing_archive_are_not_good(self):
        run=self.path/'run';run.mkdir()
        data={'run_id':'20261013-QA','hosts':['one','two'],'failed_hosts':['two'],'servers':{'one':{'counts':{'VULNERABLE':2}},'two':{'counts':{}}}}
        (run/'summary.json').write_text(json.dumps(data));source=store.import_infra(self.db,run)
        self.assertEqual(store.stats(self.db,source)['counts'],{'VULNERABLE':2})
        self.assertEqual(store.stats(self.db,source)['failed_assets'],1)
        self.assertEqual(store.stats(self.db,source)['detailed_assets'],0)
        self.assertEqual(store.stats(self.db,source)['good_rate'],0.0)
        self.assertEqual(source,store.import_infra(self.db,run));self.assertEqual(self.db.execute('SELECT COUNT(*) FROM sources').fetchone()[0],1)
        data['preflight']=True;(run/'summary.json').write_text(json.dumps(data));self.assertIsNone(store.import_infra(self.db,run))
    def test_archive_integrity_and_counts(self):
        run=self.path/'run';(run/'one').mkdir(parents=True)
        (run/'summary.json').write_text(json.dumps({'hosts':['one'],'servers':{'one':{'counts':{'GOOD':1,'VULNERABLE':1}}}}))
        with zipfile.ZipFile(run/'one'/'evidence.zip','w') as z:
            z.writestr('one/U-01.json',json.dumps({'item_id':'U-01','final_result':'GOOD'}))
            z.writestr('one/U-02.json',json.dumps({'item_id':'U-02','final_result':'VULNERABLE','importance':'상','evidence':'password=DO_NOT_LEAK'}))
        source=store.import_infra(self.db,run);stats=store.stats(self.db,source)
        self.assertEqual(stats['detailed_assets'],1);self.assertEqual(stats['good_rate'],50)
        self.assertEqual(stats['severity'],{'High':1})
        self.assertNotIn('DO_NOT_LEAK',self.db.execute('SELECT evidence FROM checks WHERE item_id="U-02"').fetchone()[0])
    def test_service_results_do_not_merge_ids_or_double_count_servers(self):
        run=self.path/'mixed';run.mkdir();data={'hosts':[],'targets':{},'servers':{}}
        for instance,profile,item,status in [('os','unix-debian','U-01','VULNERABLE'),('db-a','mysql','D-01','GOOD'),('db-b','oracle','D-01','VULNERABLE'),('was','tomcat','WEB-01','MANUAL')]:
            alias='one' if instance=='os' else 'one--'+instance;data['hosts'].append(alias)
            data['targets'][alias]={'server':'one','instance':instance,'profile':profile};data['servers'][alias]={'counts':{status:1}}
            folder=run/alias;folder.mkdir()
            with zipfile.ZipFile(folder/'evidence.zip','w') as z:z.writestr(item+'.json',json.dumps({'item_id':item,'final_result':status,'item_name':'QA '+profile}))
        (run/'summary.json').write_text(json.dumps(data));source=store.import_infra(self.db,run);stats=store.stats(self.db,source)
        self.assertEqual((stats['assets'],stats['physical_assets'],stats['physical_affected_assets']),(4,1,1))
        self.assertEqual(len(list(report.matrix.item_rows(self.db,source))),4)
        result=report.generate(self.db,{'Customer':'QA 고객사','Kind':'infra','SourceId':source},self.path)
        with zipfile.ZipFile(result['Files'][3]) as z:
            xml=z.read('xl/worksheets/sheet2.xml').decode();self.assertIn('MySQL D-01',xml);self.assertIn('Oracle D-01',xml);self.assertIn('범위 외',xml)
    def test_atomic_import_and_safe_evidence(self):
        path=self.write([{'id':'A','status':'VULNERABLE','severity':'High','evidence':{'authorization':'DO_NOT_LEAK','fact':True}}])
        source=store.import_json(self.db,path,'infra');self.assertEqual(store.stats(self.db,source)['total_checks'],1)
        self.write([{'id':'A'}, {'id':'A'}])
        with self.assertRaises(Exception):store.import_json(self.db,path,'infra')
        self.assertEqual(store.stats(self.db,source)['counts'],{'VULNERABLE':1})
        self.write([{'id':'B','images':[{'path':'../outside.png'}]}])
        with self.assertRaises(ValueError):store.import_json(self.db,path,'infra')
        path.write_text('{"connected":true}')
        with self.assertRaises(ValueError):store.import_json(self.db,path,'infra')
        self.assertEqual(store.stats(self.db,source)['total_checks'],1)
    def test_large_result_export_and_footer(self):
        rows=[{'id':'RULE-1','asset':f'server-{i:04}','title':'QA fixture finding','status':'VULNERABLE' if i%2 else 'GOOD','severity':'High','evidence':'Bearer DO_NOT_LEAK','recommendation':'QA fixture remediation'} for i in range(1200)]
        rows[1]['title']='=FORMULA'
        source=store.import_json(self.db,self.write(rows),'infra');stats=store.stats(self.db,source)
        self.assertEqual(stats['assets'],1200);self.assertEqual(stats['total_checks'],1200);self.assertEqual(stats['distinct_vulnerabilities'],1)
        req={'Customer':'QA 고객사','Kind':'infra','SourceId':source,'DetailLimit':2,'Author':'QA 담당','Period':'QA 기간','Scope':'테스트 자료만 사용'}
        result=report.generate(self.db,req,self.path)
        doc=Document(result['Files'][0]);self.assertIn('QA 고객사',doc.sections[0].footer.paragraphs[0].text)
        self.assertIn('PAGE',doc.sections[0].footer._element.xml);self.assertIn('NUMPAGES',doc.sections[0].footer._element.xml)
        self.assertIn('인프라 서버 및 네트워크',doc.core_properties.title)
        with open(result['Files'][1],encoding='utf-8-sig',newline='') as f:records=list(csv.reader(f))
        self.assertEqual(len(records),1201);self.assertTrue(any(r[2]=="'=FORMULA" for r in records[1:]))
        self.assertNotIn('DO_NOT_LEAK',str(records))
        self.assertEqual(len([p for p in doc.paragraphs if 'QA fixture finding' in p.text]),0)
        self.assertIn('점검 결과 보고서',doc.core_properties.title)
        self.assertEqual(len(doc.tables[-1].rows),1201)
        with zipfile.ZipFile(result['Files'][3]) as z:
            sheet=ET.fromstring(z.read('xl/worksheets/sheet2.xml'))
        ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        matrix_rows=[[''.join(c.itertext()) for c in row] for row in sheet.findall('m:sheetData/m:row',ns)]
        self.assertEqual(len(matrix_rows),1201);self.assertEqual(matrix_rows[1][4],'양호');self.assertEqual(matrix_rows[2][4],'취약')
        self.assertEqual(matrix_rows[-1][0],'server-1199')
        self.assertEqual(sheet.find('m:sheetViews/m:sheetView/m:pane',ns).attrib['topLeftCell'],'E2')
        with self.assertRaises(ValueError):report.generate(self.db,dict(req,Kind='web'),self.path)
    def test_templates_do_not_fabricate_results(self):
        for kind in ('mobile','web','infra'):
            result=report.generate(self.db,{'Customer':'고객사명','Kind':kind},self.path,True)
            self.assertIsNone(result['Stats']);self.assertTrue(result['Draft'])
            doc=Document(result['Files'][0]);self.assertIn('검토용 서식',[p.text for p in doc.paragraphs])
            self.assertTrue(any(c.text=='미입력' for t in doc.tables for r in t.rows for c in r.cells))
    def test_evidence_image_integrity(self):
        image=self.path/'evidence.png';Image.new('RGB',(480,240),'white').save(image)
        source=store.import_json(self.db,self.write([{'id':'QA-image','status':'VULNERABLE','evidence':{'command':'QA command','output':'QA output'},'images':[{'path':'evidence.png','caption':'QA image'}]}],kind='mobile'),'mobile')
        evidence=self.db.execute('SELECT evidence FROM checks WHERE source_id=?',(source,)).fetchone()[0]
        self.assertIn('확인 결과',evidence);self.assertNotIn('{',evidence)
        req={'Customer':'QA 고객사','Kind':'mobile','SourceId':source,'IncludeImages':True}
        result=report.generate(self.db,req,self.path);doc=Document(result['Files'][0]);self.assertEqual(len(doc.inline_shapes),2)
        Image.new('RGB',(480,240),'black').save(image)
        result=report.generate(self.db,req,self.path);doc=Document(result['Files'][0]);self.assertEqual(len(doc.inline_shapes),1)
        self.assertTrue(any('변경되어' in p.text for p in doc.paragraphs))

if __name__=='__main__':unittest.main(verbosity=2)
